"""Compare the V1 and V2 EXECUTE reports and explain every PDB membership change."""

from __future__ import annotations

import argparse
import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def by_pdb(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    result = {row["PDB_ID"]: row for row in rows}
    assert len(result) == len(rows), "Duplicate PDB_ID in a report"
    return result


def as_number(value: str) -> float | None:
    return float(value) if value else None


def explain_added(old: dict[str, str], audit: dict[str, str]) -> str:
    status = old["FINAL_STATUS"]
    if status == "EXCLUDE_TRAIN_RMSD_CUTOFF":
        rmsd = old["STRICT_RANK1_RMSD_ANGSTROM"]
        assert rmsd and float(rmsd) > 15
        return f"旧版按 rank-1 RMSD >15 Å 排除了整个 PDB（旧值 {rmsd} Å）；新版仅记录该值，不据此排除。"
    if status == "EXCLUDE_FROZEN_RMSD":
        return (
            "旧版冻结 RMSD 排除清单使该 PDB 未入选；新版不沿用这项 RMSD 排除。"
            f"旧版 RMSD 评估状态：{audit['RMSD_EVAL_STATUS'] or '未记录'}。"
            "此状态本身不表示 Protenix 预测失败。"
        )
    if status == "EXCLUDE_NO_STRICT_RANK1":
        return "旧版因缺少有效 strict rank-1 RMSD 记录而未入选；新版允许保留，并仅作 rank-1 注释。"
    raise ValueError(f"Unexpected prior status for added PDB: {status}")


def style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def build(old_dir: Path, new_dir: Path, output: Path) -> None:
    old_manifest = by_pdb(read_tsv(old_dir / "final_manifest.tsv"))
    new_manifest = by_pdb(read_tsv(new_dir / "final_manifest.tsv"))
    old_audit = by_pdb(read_tsv(old_dir / "selection_audit.tsv"))
    new_audit = by_pdb(read_tsv(new_dir / "selection_audit.tsv"))
    old_source = by_pdb(read_tsv(old_dir / "source_inventory.tsv"))
    new_source = by_pdb(read_tsv(new_dir / "source_inventory.tsv"))
    chain_rows = by_pdb(read_tsv(new_dir / "test_chain_evaluation.tsv"))
    internal_hits = read_tsv(new_dir / "test_internal_hits.tsv")

    # Both reports must describe the same original CIF set before membership is compared.
    assert old_source == new_source, "The two runs do not share the same CIF inventory and SHA256 values"
    assert old_manifest.keys() == new_manifest.keys()
    assert old_manifest.keys() <= old_source.keys()
    old_selected = {p: row["FINAL_SPLIT"] for p, row in old_manifest.items() if row["FINAL_SPLIT"]}
    new_selected = {p: row["FINAL_SPLIT"] for p, row in new_manifest.items() if row["FINAL_SPLIT"]}
    removed = sorted(old_selected.keys() - new_selected.keys())
    added = sorted(new_selected.keys() - old_selected.keys())
    retained = old_selected.keys() & new_selected.keys()
    moved = [p for p in retained if old_selected[p] != new_selected[p]]
    assert not moved, f"PDBs changed split: {moved}"
    assert (len(removed), len(added)) == (24, 51)

    wb = Workbook()
    overview = wb.active
    overview.title = "概览"
    overview.append(["项目", "旧版 V1", "新版 V2", "说明"])
    overview.append(["报告来源", old_dir.name, new_dir.name, "均为 EXECUTE 日志"])
    overview.append(["原始 CIF 数", len(old_source), len(new_source), "按 PDB_ID 与 SHA256 核对一致"])
    for split, label in (("train", "训练集"), ("val", "验证集"), ("test", "测试集")):
        n_old = sum(v == split for v in old_selected.values())
        n_new = sum(v == split for v in new_selected.values())
        overview.append([label, n_old, n_new, f"净变化 {n_new - n_old:+d}"])
    overview.append(["已入选 PDB 总数", len(old_selected), len(new_selected), f"净增加 {len(added) - len(removed)}"])
    overview.append(["从旧版移出", len(removed), "", "见“移出24个”工作表；均原属测试集"])
    overview.append(["新版新增", "", len(added), "见“新增51个”工作表"])
    overview.append(["跨 train/val/test 移动", 0, 0, "保留条目无跨集合移动"])
    overview.append(["比较范围", "PDB_ID", "PDB_ID", "比较最终入选名单；原始 CIF 未删除"])
    overview.append(["移出：与 train/val 同源", "", 22, "短序列全局比对命中；具体对象和指标见明细"])
    overview.append(["移出：test 内部冗余", "", 2, "保留同源组代表，移出非代表条目"])
    reason_counts = Counter(old_manifest[p]["FINAL_STATUS"] for p in added)
    for status, label in (
        ("EXCLUDE_TRAIN_RMSD_CUTOFF", "新增：旧版 rank-1 RMSD >15 Å"),
        ("EXCLUDE_FROZEN_RMSD", "新增：旧版冻结 RMSD 清单"),
        ("EXCLUDE_NO_STRICT_RANK1", "新增：旧版缺 rank-1 记录"),
    ):
        overview.append([label, "", reason_counts[status], status])
    overview.append(["同源判据", "", "", "序列一致性 ≥0.8，查询和目标覆盖率均 ≥0.8；详见 V2 报告参数"])
    style_sheet(overview, [34, 61, 61, 78])

    removed_sheet = wb.create_sheet("移出24个")
    removed_sheet.append([
        "PDB_ID", "旧版划分", "新版结果", "具体原因", "匹配对象或保留代表 PDB",
        "匹配对象在旧版的划分", "序列一致性", "查询覆盖率", "对象覆盖率", "比对来源",
        "RNA 长度 nt", "原始 CIF 路径", "原始 CIF SHA256", "旧版目录",
    ])
    removed_reasons = Counter()
    for pdb in removed:
        old = old_manifest[pdb]
        new = new_manifest[pdb]
        chain = chain_rows[pdb]
        status = new["FINAL_STATUS"]
        assert old_selected[pdb] == "test"
        if status == "DROP_REFERENCE_HOMOLOG":
            target = chain["MATCH_PDB_ID"]
            identity = as_number(chain["IDENTITY"])
            query_cov = as_number(chain["QUERY_COVERAGE"])
            target_cov = as_number(chain["TARGET_COVERAGE"])
            source = chain["ALIGNMENT_SOURCE"]
            assert target in old_selected and old_selected[target] in {"train", "val"}
            assert source == "SHORT_GLOBAL_FALLBACK"
            assert min(identity, query_cov, target_cov) >= 0.8
            reason = f"与旧版 {old_selected[target]} 的 {target} 同源；新版测试集去同源时移出。"
        elif status == "DROP_INTERNAL_REDUNDANT":
            target = chain["REPRESENTATIVE_PDB_ID"]
            assert target in new_selected and new_selected[target] == "test"
            matches = [h for h in internal_hits if {h["QUERY_PDB_ID"], h["TARGET_PDB_ID"]} == {pdb, target}]
            assert matches, f"Missing internal homology hit for {pdb} and {target}"
            hit = matches[0]
            identity = as_number(hit["IDENTITY"])
            if hit["QUERY_PDB_ID"] == pdb:
                query_cov = as_number(hit["QUERY_COVERAGE"])
                target_cov = as_number(hit["TARGET_COVERAGE"])
            else:
                query_cov = as_number(hit["TARGET_COVERAGE"])
                target_cov = as_number(hit["QUERY_COVERAGE"])
            source = hit["ALIGNMENT_SOURCE"]
            assert min(identity, query_cov, target_cov) >= 0.8
            reason = f"测试集内部与 {target} 同源；{target} 被保留为代表，本条目移出。"
        else:
            raise ValueError(f"Unexpected removal status for {pdb}: {status}")
        removed_reasons[status] += 1
        removed_sheet.append([
            pdb, old_selected[pdb], status, reason, target, old_selected.get(target, "旧版未入选"),
            identity, query_cov, target_cov, source, as_number(new_audit[pdb]["RNA_LENGTH"]),
            old_source[pdb]["CIF_PATH"], old_source[pdb]["SHA256"], old["TARGET_DIRECTORY"],
        ])
        for col in (7, 8, 9):
            removed_sheet.cell(removed_sheet.max_row, col).number_format = "0.000000"
    assert removed_reasons == {"DROP_REFERENCE_HOMOLOG": 22, "DROP_INTERNAL_REDUNDANT": 2}
    style_sheet(removed_sheet, [13, 14, 29, 66, 24, 24, 18, 18, 18, 29, 16, 78, 66, 78])

    added_sheet = wb.create_sheet("新增51个")
    added_sheet.append([
        "PDB_ID", "新版划分", "旧版未入选状态", "旧版未入选具体原因", "旧版 rank-1 RMSD Å",
        "旧版 RMSD 评估状态", "新版 RMSD 评估状态", "RNA 长度 nt", "原始 CIF 路径",
        "原始 CIF SHA256", "新版目录",
    ])
    for pdb in added:
        old = old_manifest[pdb]
        new = new_manifest[pdb]
        audit = old_audit[pdb]
        assert new["FINAL_STATUS"] == "KEPT"
        added_sheet.append([
            pdb, new_selected[pdb], old["FINAL_STATUS"], explain_added(old, audit),
            as_number(old["STRICT_RANK1_RMSD_ANGSTROM"]), audit["RMSD_EVAL_STATUS"],
            new_audit[pdb]["RMSD_EVAL_STATUS"], as_number(new_audit[pdb]["RNA_LENGTH"]),
            old_source[pdb]["CIF_PATH"], old_source[pdb]["SHA256"], new["TARGET_DIRECTORY"],
        ])
        added_sheet.cell(added_sheet.max_row, 5).number_format = "0.000"
    assert reason_counts == {
        "EXCLUDE_TRAIN_RMSD_CUTOFF": 28,
        "EXCLUDE_FROZEN_RMSD": 12,
        "EXCLUDE_NO_STRICT_RANK1": 11,
    }
    style_sheet(added_sheet, [13, 14, 31, 85, 22, 25, 25, 17, 78, 66, 78])

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    checked = load_workbook(output, read_only=True)
    assert checked.sheetnames == ["概览", "移出24个", "新增51个"]
    assert checked["移出24个"].max_row == 25
    assert checked["新增51个"].max_row == 52
    removed_ids = {row[0] for row in checked["移出24个"].iter_rows(min_row=2, values_only=True)}
    added_ids = {row[0] for row in checked["新增51个"].iter_rows(min_row=2, values_only=True)}
    assert removed_ids == set(removed)
    assert added_ids == set(added)
    assert all(row[3] for row in checked["移出24个"].iter_rows(min_row=2, values_only=True))
    assert all(row[3] for row in checked["新增51个"].iter_rows(min_row=2, values_only=True))
    print(f"Created {output}")
    print(f"Old {len(old_selected)}, new {len(new_selected)}, removed {len(removed)}, added {len(added)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--old-report", required=True, type=Path)
    parser.add_argument("--new-report", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(args.old_report, args.new_report, args.output)
